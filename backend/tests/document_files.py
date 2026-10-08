"""Synthetic CV files built in memory for tests (no real CV is ever used)."""

import io
import zipfile

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

CV_LINES = [
    "Alex Example",
    "Senior Software Engineer",
    "Email: alex.example@example.test  Phone: +351 912 345 678",
    "Profile: https://www.linkedin.com/in/alex-example",
    "Experience: Northwind Labs 2019-03 to 2023-11, Staff Engineer. Built Python and PostgreSQL services.",
    "Skills: Python, PostgreSQL, TypeScript, Docker",
]


def pdf_bytes(lines: list[str] | None = None, pages: int = 1) -> bytes:
    """A text PDF; ``lines`` empty gives pages with no text (like a scan)."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    for _ in range(pages):
        y = 800
        for line in CV_LINES if lines is None else lines:
            pdf.drawString(50, y, line)
            y -= 18
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def docx_bytes(lines: list[str] | None = None, table: list[list[str]] | None = None) -> bytes:
    """A Word document with paragraphs and an optional table."""
    document = DocxDocument()
    for line in CV_LINES if lines is None else lines:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def text_bytes(lines: list[str] | None = None, encoding: str = "utf-8") -> bytes:
    """A plain-text CV."""
    return "\n".join(CV_LINES if lines is None else lines).encode(encoding)


def zip_bomb_docx(expanded_mb: int = 120) -> bytes:
    """A .docx whose parts expand far beyond the limit while the file itself stays small."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b"0" * (expanded_mb * 1024 * 1024))
    return buffer.getvalue()
